
import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BT = chr(96)
FENCE = BT * 3

SRC = sys.argv[1]
OUT = sys.argv[2]

doc = Document()
st = doc.styles["Normal"]
st.font.name = "宋体"; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
st.paragraph_format.line_spacing = 1.5
st.paragraph_format.space_after = Pt(4)

for name, size in (("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11.5)):
    hs = doc.styles[name]
    hs.font.name = "黑体"; hs.font.size = Pt(size)
    hs.font.bold = True
    hs.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    hs.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    hs.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 9)
    hs.paragraph_format.space_after = Pt(6 if name == "Heading 1" else 4)

sec = doc.sections[0]
sec.top_margin = Cm(2.3); sec.bottom_margin = Cm(2.3)
sec.left_margin = Cm(2.6); sec.right_margin = Cm(2.6)

fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
rr = OxmlElement("w:r"); tt = OxmlElement("w:t"); tt.text = "1"
rr.append(tt); fld.append(rr); fp._p.append(fld)

def cn(run, name):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)

def add_runs(p, text, base_font="宋体", size=10.5):
    text = text.replace(BT, "")
    for seg in re.split(r"(\*\*.+?\*\*)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            r = p.add_run(seg[2:-2]); r.bold = True
        else:
            r = p.add_run(seg)
        cn(r, base_font); r.font.size = Pt(size)

lines = open(SRC, encoding="utf-8").read().split("\n")

meta = {}
body = []
for ln in lines:
    m = re.match(r"^-\s+\*\*(.+?)\*\*：(.+)$", ln.strip())
    if m and not body:
        meta[m.group(1)] = m.group(2).strip()
    elif ln.strip().startswith("# ") and not body:
        title = ln.strip()[2:]
        body.append(ln); meta["__title__"] = title
    else:
        body.append(ln)

# ---------- 封面 ----------
for _ in range(2):
    doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("慢性病患者健康风险预测与干预"); cn(r, "黑体"); r.font.size = Pt(30); r.bold = True

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(18)
r = p.add_run("方 案 说 明"); cn(r, "黑体"); r.font.size = Pt(20)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(10)
r = p.add_run("目标 · 主要思路 · 实施步骤"); cn(r, "宋体"); r.font.size = Pt(12)
r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

for _ in range(6):
    doc.add_paragraph()

rows = [
    ("学生姓名", "许宇翔"),
    ("学　　号", "20240517108"),
    ("课程名称", "海量数据挖掘分析与可视化实战（综合项目）"),
    ("文档类型", "M2 里程碑交付物 · 技术方案与架构设计"),
    ("数据来源", "CDC BRFSS LLCP 2019–2023 五年合并，2,137,374 条原始记录"),
    ("文档版本", "v1.0"),
]
t = doc.add_table(rows=0, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for k, v in rows:
    cells = t.add_row().cells
    cells[0].width = Cm(3.4); cells[1].width = Cm(10.6)
    for idx, txt in enumerate((k, v)):
        cp = cells[idx].paragraphs[0]
        cp.paragraph_format.space_after = Pt(2)
        rr2 = cp.add_run(txt); cn(rr2, "宋体"); rr2.font.size = Pt(11)
        if idx == 0:
            rr2.bold = True

doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("2025 年"); cn(r, "宋体"); r.font.size = Pt(12)

doc.add_page_break()

# ---------- 目录 ----------
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(12)
r = p.add_run("目　　录"); cn(r, "黑体"); r.font.size = Pt(16); r.bold = True

for ln in body:
    m = re.match(r"^##\s+(.*)$", ln)
    if m:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_after = Pt(5)
        cp.paragraph_format.left_indent = Cm(0.6)
        txt = m.group(1).replace(BT, "")
        r = cp.add_run(txt); cn(r, "宋体"); r.font.size = Pt(11)

doc.add_page_break()

# ---------- 正文 ----------
i = 0
n_tab = n_code = n_head = 0
while i < len(body):
    ln = body[i]
    if ln.strip() == "---" or ln.strip() == "":
        i += 1; continue
    if ln.startswith("|") and i + 1 < len(body) and re.match(r"^\|[\s:\-|]+\|$", body[i+1].strip()):
        rows2 = []
        while i < len(body) and body[i].startswith("|"):
            cells = [c.strip() for c in body[i].strip().strip("|").split("|")]
            if not re.match(r"^[\s:\-]+$", "".join(cells)):
                rows2.append(cells)
            i += 1
        ncol = max(len(r) for r in rows2)
        tb = doc.add_table(rows=0, cols=ncol)
        tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, row in enumerate(rows2):
            cs = tb.add_row().cells
            for ci in range(ncol):
                cell = cs[ci]; cell.text = ""
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_after = Pt(1)
                cp.paragraph_format.line_spacing = 1.15
                add_runs(cp, (row[ci] if ci < len(row) else ""), size=9)
                if ri == 0:
                    for r in cp.runs:
                        r.bold = True
        doc.add_paragraph()
        n_tab += 1
        continue
    if ln.strip().startswith(FENCE):
        i += 1; buf = []
        while i < len(body) and not body[i].strip().startswith(FENCE):
            buf.append(body[i]); i += 1
        i += 1
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after = Pt(8)
        cp.paragraph_format.line_spacing = 1.0
        r = cp.add_run("\n".join(buf)); cn(r, "Consolas"); r.font.size = Pt(7.5)
        n_code += 1
        continue
    m = re.match(r"^(#{1,4})\s+(.*)$", ln)
    if m:
        lvl = len(m.group(1)); txt = m.group(2).replace(BT, "")
        if lvl == 1:
            i += 1; continue
        style = "Heading %d" % (lvl - 1) if lvl <= 3 else "Heading 3"
        hp = doc.add_paragraph(style=style)
        r = hp.add_run(txt); cn(r, "黑体")
        r.font.size = Pt(15 if lvl == 2 else 13 if lvl == 3 else 11.5)
        r.bold = True
        n_head += 1
        i += 1; continue
    m = re.match(r"^(\d+)\.\s+(.*)$", ln.strip())
    if m:
        cp = doc.add_paragraph()
        cp.paragraph_format.left_indent = Cm(0.74)
        cp.paragraph_format.first_line_indent = Cm(-0.74)
        r = cp.add_run(m.group(1) + ". "); cn(r, "宋体"); r.font.size = Pt(10.5)
        add_runs(cp, m.group(2))
        i += 1; continue
    if ln.strip().startswith("- "):
        cp = doc.add_paragraph()
        cp.paragraph_format.left_indent = Cm(0.74)
        cp.paragraph_format.first_line_indent = Cm(-0.4)
        r = cp.add_run("・"); cn(r, "宋体"); r.font.size = Pt(10.5)
        add_runs(cp, ln.strip()[2:])
        i += 1; continue
    cp = doc.add_paragraph()
    add_runs(cp, ln.strip())
    i += 1

doc.save(OUT)
print("tables=%d code=%d headings=%d" % (n_tab, n_code, n_head))
