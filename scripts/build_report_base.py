# -*- coding: utf-8 -*-
"""build_report.py — 按《大数据分析报告模板.docx》逐章生成课程设计报告。

所有数字来自 artifacts/metrics/*.json 与真实运行结果，不手抄。
"""
import json
import os
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJ = os.path.dirname(SCRIPT_DIR)          # 项目根
ROOT = os.path.dirname(PROJ)                # 上级目录（最终 docx 落地处）
MET = os.path.join(PROJ, "artifacts", "metrics")
FIG = os.path.join(PROJ, "artifacts", "figures")
SHOT = os.path.join(PROJ, "web", "assets", "screenshots")
OUTDOCX = os.path.join(ROOT, "大数据分析课程设计报告.docx")

KECHENG = "大数据分析课程设计"
XIANGMU = "基于 BRFSS 大数据的慢性病患者健康风险预测与干预分析"
ZUOZHE = "许宇翔"
XUEHAO = "20240517108"
RENWU = ("独立完成全链路：数据获取与原始数据留存、PySpark 数据清洗与建库、探索性分析与假设检验、"
         "预测建模与对比实验、FastAPI 后端与 ECharts 可视化、工程化与文档撰写")


def jload(name):
    p = os.path.join(MET, name)
    with open(p, encoding="utf-8") as f:
        return json.load(f)


class Rep:
    def __init__(self):
        self.doc = Document()
        self._setup()

    def _setup(self):
        st = self.doc.styles["Normal"]
        st.font.name = "宋体"
        st.font.size = Pt(10.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        for sec in self.doc.sections:
            sec.top_margin = Cm(2.3); sec.bottom_margin = Cm(2.3)
            sec.left_margin = Cm(2.6); sec.right_margin = Cm(2.4)

    def _cn(self, run, font):
        run.font.name = font
        run.element.rPr.rFonts.set(qn("w:eastAsia"), font)

    def h(self, text, level):
        p = self.doc.add_heading(level=level)
        r = p.add_run(text)
        self._cn(r, "黑体")
        r.font.color.rgb = RGBColor(0, 0, 0)
        r.font.size = Pt({0: 22, 1: 16, 2: 13.5, 3: 12}[level])
        r.bold = True
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(6)
        return p

    def p(self, text, indent=True, size=10.5, bold=False, color=None):
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(21) if indent else None
        p.paragraph_format.line_spacing = 1.45
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        self._cn(r, "宋体")
        r.font.size = Pt(size)
        r.bold = bold
        if color:
            r.font.color.rgb = color
        return p

    def bullets(self, items):
        for it in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing = 1.35
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(it)
            self._cn(r, "宋体")
            r.font.size = Pt(10.5)

    def table(self, cols, rows, widths=None, caption=None):
        if caption:
            cp = self.doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cp.add_run(caption)
            self._cn(r, "宋体"); r.font.size = Pt(9); r.bold = True
        t = self.doc.add_table(rows=1, cols=len(cols))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = t.rows[0].cells
        for i, c in enumerate(cols):
            hdr[i].text = ""
            r = hdr[i].paragraphs[0].add_run(str(c))
            self._cn(r, "黑体"); r.font.size = Pt(9); r.bold = True
            hdr[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run("" if v is None else str(v))
                self._cn(r, "宋体"); r.font.size = Pt(9)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def fig(self, path, caption, width_cm=15.0):
        if not os.path.exists(path):
            self.p("[缺失图件] %s" % os.path.basename(path), indent=False)
            return
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.add_run().add_picture(path, width=Cm(width_cm))
        cp = self.doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(caption)
        self._cn(r, "宋体"); r.font.size = Pt(9)
        cp.paragraph_format.space_after = Pt(8)

    def code(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        r.font.name = "Consolas"
        r.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        r.font.size = Pt(8.5)

    def pagebreak(self):
        self.doc.add_page_break()

    def save(self, path):
        self.doc.save(path)
