
import re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BT = chr(96)
FENCE = BT * 3

SRC = r"docs/技术方案与架构设计文档.md"
OUT = r"docs/技术方案与架构设计文档.docx"

doc = Document()
st = doc.styles["Normal"]
st.font.name = "宋体"; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
st.paragraph_format.line_spacing = 1.45
st.paragraph_format.space_after = Pt(4)

for s in doc.sections:
    s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.2)
    s.left_margin = Cm(2.6); s.right_margin = Cm(2.6)

def cn(run, name):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)

def add_runs(p, text, base_font="宋体", size=10.5):
    text = text.replace(BT, "")
    for seg in re.split(r"(\*\*.+?\*\*)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            r = p.add_run(seg[2:-2]); r.bold = True
        else:
            r = p.add_run(seg)
        cn(r, base_font); r.font.size = Pt(size)

lines = open(SRC, encoding="utf-8").read().split("\n")
i = 0
n_tab = n_code = n_head = 0
while i < len(lines):
    ln = lines[i]
    if ln.strip() == "---":
        i += 1; continue
    if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|$", lines[i+1].strip()):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not re.match(r"^[\s:\-]+$", "".join(cells)):
                rows.append(cells)
            i += 1
        ncol = max(len(r) for r in rows)
        t = doc.add_table(rows=0, cols=ncol)
        t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for ci in range(ncol):
                txt = row[ci] if ci < len(row) else ""
                cell = cells[ci]
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.15
                add_runs(p, txt, size=9)
                if ri == 0:
                    for r in p.runs:
                        r.bold = True
        doc.add_paragraph()
        n_tab += 1
        continue
    if ln.strip().startswith(FENCE):
        i += 1; buf = []
        while i < len(lines) and not lines[i].strip().startswith(FENCE):
            buf.append(lines[i]); i += 1
        i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run("\n".join(buf))
        cn(r, "Consolas"); r.font.size = Pt(7.5)
        n_code += 1
        continue
    m = re.match(r"^(#{1,4})\s+(.*)$", ln)
    if m:
        lvl = len(m.group(1)); txt = m.group(2).replace(BT, "")
        if lvl == 1:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(14)
            r = p.add_run(txt); cn(r, "黑体"); r.font.size = Pt(20); r.bold = True
        elif lvl == 2:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(6)
            r = p.add_run(txt); cn(r, "黑体"); r.font.size = Pt(14); r.bold = True
        elif lvl == 3:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(9); p.paragraph_format.space_after = Pt(4)
            r = p.add_run(txt); cn(r, "黑体"); r.font.size = Pt(12); r.bold = True
        else:
            p = doc.add_paragraph()
            r = p.add_run(txt); cn(r, "黑体"); r.font.size = Pt(10.5); r.bold = True
        n_head += 1
        i += 1; continue
    m = re.match(r"^(\d+)\.\s+(.*)$", ln.strip())
    if m:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.74)
        r = p.add_run(m.group(1) + ". "); cn(r, "宋体"); r.font.size = Pt(10.5)
        add_runs(p, m.group(2))
        i += 1; continue
    if ln.strip().startswith("- "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        r = p.add_run("・"); cn(r, "宋体"); r.font.size = Pt(10.5)
        add_runs(p, ln.strip()[2:])
        i += 1; continue
    if ln.strip() == "":
        i += 1; continue
    p = doc.add_paragraph()
    add_runs(p, ln.strip())
    i += 1

doc.save(OUT)
print("tables=%d code=%d headings=%d" % (n_tab, n_code, n_head))
print("saved", OUT)
